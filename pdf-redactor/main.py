## Warning filter
import warnings

warnings.filterwarnings(
    "ignore",
    message="The `fitz` API is deprecated.*",
    category=FutureWarning
)

##re for regex, and spacy for NER models
import argparse
from pathlib import Path
from pdf2docx import Converter
import re
import spacy
from docx import Document
import tempfile
import pymupdf
import shlex
import json

parser = argparse.ArgumentParser(description="Convert a PDF to DOCX and mask sensitive information")

##load spacCy's small English model
nlp = spacy.load("en_core_web_sm")

placeholder_map = {}

def save_redaction_key(key_path):
    key_document = Document();

    key_document.add_heading("Redaction Key", level=1)

    key_document.add_paragraph(
        "This document contains the original sensitive information "
        "that was replaced during redaction. Keep it private."
    )

    table = key_document.add_table(rows=1, cols=2)
    table.style = "Table Grid"

    header = table.rows[0].cells
    header[0].text = "Placeholder"
    header[1].text = "Original Value"

    for (category, original_value), placeholder in placeholder_map.items():
        row = table.add_row().cells
        row[0].text = placeholder
        row[1].text = original_value

    key_document.save(key_path)

def get_placeholder(category, value):
    key = (category, value)

    if key in placeholder_map:
        return placeholder_map[key]

    count = sum(
        1 for stored_category, _ in placeholder_map
        if stored_category == category
    )

    placeholder = f"[{category}_{count+1}]"
    placeholder_map[key] = placeholder

    return placeholder


def mask_paragraph(paragraph):
    runs = paragraph.runs

    if not runs:
        return

    original_texts = [run.text for run in runs]
    full_text = "".join(original_texts)

    spans = find_sensitive_spans(full_text)

    if not spans:
        return

    # Determine each run's position in original paragraph
    position = 0

    # Construct each run's new text without shifting positions
    for i, run in enumerate(runs):
        original = original_texts[i]
        
        run_start = position
        run_end = position + len(original)
        position = run_end

        ## Find sensitive spans overlapping this run
        relevant = [
            (start, end, placeholder)
            for start, end, placeholder in spans
            if start < run_end and end > run_start
        ]

        ## Work right to left to avoid shifting indices
        for start, end, placeholder in reversed(relevant):
            local_start = max(0, start - run_start)
            local_end = min(len(original_texts[i]), end - run_start)

            # Insert placeholder only where the match begins
            replacement = (
                placeholder if run_start <= start < run_end else ""
            )

            original = (
                original[:local_start] + replacement + original[local_end:]
            )

        if relevant:
            run.text = original

    
def find_sensitive_spans(text):
    spans = []
    patterns = {
        "EMAIL": r"\b[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
        "PHONE": r"\b\d{3}[-.]?\d{3}[-.]?\d{4}\b",
        "SSN": r"\b\d{3}-\d{2}-\d{4}\b"
    }

    for category, pattern in patterns.items():
        for match in re.finditer(pattern, text):
            spans.append((
                match.start(),
                match.end(),
                category
            ))

    normalized_text = text.replace("\u200b", " ")
    doc = nlp(normalized_text)
    excluded_names = {
        "email",
        "name",
        "phone",
        "number",
        "address",
        "contact",
        "ssn"
    }

    for entity in doc.ents:
        if (entity.label_ != "PERSON"):
            continue;

        start = entity.start_char
        end = entity.end_char

        #Remove surrounding whitepace and separator characters
        while start < end and text[start] in " \t\n\u200b-":
            start += 1

        while end > start and text[end-1] in " \t\n\u200b-":
            end -= 1

        name = text[start:end].replace("\u200b", " ").strip()

        if name and name.lower() not in excluded_names:
            spans.append((start, end, "PERSON"))

    spans.sort(key=lambda span: (span[0], -(span[1] - span[0])))

    filtered = []

    for span in spans:
        if not filtered or span[0] >= filtered[-1][1]:
            filtered.append(span)

    numbered_spans = []

    for start, end, category in filtered:
        original_value = text[start:end]
        placeholder = get_placeholder(category, original_value)
        numbered_spans.append((start, end, placeholder))

    return numbered_spans

#Account for containers like table/header/footer
def mask_container(container):
    for paragraph in container.paragraphs:
        mask_paragraph(paragraph)
        repair_spacing(paragraph)
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                mask_container(cell)


#Fix spacing
def repair_spacing(paragraph):
    runs = paragraph.runs

    #Record each character and the run it belongs to
    chars = [
        (char, i)
        for i, run in enumerate(runs)
        for char in run.text
    ]

    updated = ["" for _ in runs]
    pos = 0

    while pos < len(chars):
        char, run_index = chars[pos]

        if char != "\u200b":
            updated[run_index] += char
            pos += 1
            continue
    

        end = pos
        while end < len(chars) and chars[end][0] == "\u200b":
            end += 1

        left = chars[pos - 1][0] if pos else ""
        right = chars[end][0] if end < len(chars) else ""

        #Decide whether to insert a real space
        if (left and right and not left.isspace() and not right.isspace()
            and right not in ".,;:!?)]}" and left not in "({["
        ):
            updated[run_index] += " "

        pos = end

    #Update existing runs to preserve formatting
    for run, text in zip(runs, updated):
        if run.text != text:
            run.text = text

parser.add_argument(
    "-p",
    "--pdf",
    required=True,
    help="Path to the PDF file"
)

def remove_pdf_links(pdf_path, temp_path):
    with pymupdf.open(str(pdf_path)) as pdf:
        for page in pdf:
            for link in page.get_links():
                page.delete_link(link)

        pdf.save(str(temp_path))

args = parser.parse_args();

pdf_path = Path(args.pdf).expanduser().resolve()
if not pdf_path.is_file():
    parser.error("The specified file does not exist.")

if pdf_path.suffix.lower() != ".pdf":
    parser.error("The specified file must be a PDF.")

print("Valid PDF file:", pdf_path)

file_name = pdf_path.stem
output_name = file_name + "_redacted.docx"
output_path = pdf_path.with_name(output_name)

print("Input:", pdf_path)
print("Output:", output_path)

##temporary directory to remove links
with tempfile.TemporaryDirectory() as temp_dir:
    temp_path = Path(temp_dir) / "no_links.pdf"
    remove_pdf_links(pdf_path, temp_path)
    converter = Converter(str(temp_path))

    try:
        converter.convert(str(output_path))
    finally:
        converter.close()

document = Document(output_path)

mask_container(document)

for section in document.sections:
    mask_container(section.header)
    mask_container(section.footer)

document.save(output_path)

key_path = pdf_path.with_name(pdf_path.stem + "_redacted_KEY.docx")
save_redaction_key(key_path)

print("Conversion and masking successful!")
print(f"open {shlex.quote(str(output_path))}")
print(f"Open redaction key: open{shlex.quote(str(key_path))}")


